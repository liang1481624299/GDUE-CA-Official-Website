"""成员管理接口（现任/往届、归档、头像上传）。

公开接口：GET /api/members 按 display_order 升序，过滤归档项。
管理接口：CRUD + 头像上传（5MB 限制 + 仅图片 + 自动压缩）。
"""
import io
import uuid
from pathlib import Path
from typing import Literal

from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, Request, Response, UploadFile, status
from PIL import Image
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.middleware import get_client_ip
from app.core.permissions import require_permission
from app.db.models import AuditLog, MediaCategory, MediaFile, Member, MemberTerm
from app.db.session import get_db
from app.schemas.member import (
    AvatarUploadOut,
    MemberAdminOut,
    MemberBatchUpdate,
    MemberCreate,
    MemberImportError,
    MemberImportOut,
    MemberOut,
    MemberUpdate,
)
from app.utils.crud import apply_eq, apply_search, paginate
from app.utils.export import to_xlsx

public_router = APIRouter(prefix="/api/members", tags=["members-public"])
admin_router = APIRouter(prefix="/api/admin/members", tags=["members-admin"])

MEMBER_AVATAR_DIR = Path("uploads/members")
MEMBER_AVATAR_MAX_SIZE = 5 * 1024 * 1024  # 5MB


def _log(db: AsyncSession, uid: int, action: str, target: str | None, ip: str | None):
    db.add(AuditLog(user_id=uid, action=action, target=target, ip=ip))


def _sniff_image(data: bytes) -> str | None:
    """按文件头魔数识别图片类型。"""
    if data.startswith(b"\xff\xd8\xff"):
        return "JPEG"
    if data.startswith(b"\x89PNG\r\n\x1a\n"):
        return "PNG"
    if len(data) >= 12 and data[:4] == b"RIFF" and data[8:12] == b"WEBP":
        return "WEBP"
    return None


def _compress_avatar(data: bytes, fmt: str) -> tuple[bytes, str, str]:
    """头像压缩：限制最长边 512px，统一输出 JPEG quality=85（体积小，预览清晰）。

    返回 (压缩后字节, 扩展名, mime)。
    """
    img = Image.open(io.BytesIO(data))
    # 转 RGB（PNG 带 alpha 通道时无法直接存 JPEG）
    if img.mode in ("RGBA", "LA", "P"):
        img = img.convert("RGB")
    # 最长边 512px
    max_side = 512
    if max(img.size) > max_side:
        ratio = max_side / max(img.size)
        img = img.resize(
            (int(img.size[0] * ratio), int(img.size[1] * ratio)),
            Image.Resampling.LANCZOS,
        )
    buf = io.BytesIO()
    img.save(buf, format="JPEG", quality=85, optimize=True)
    return buf.getvalue(), ".jpg", "image/jpeg"


# ---------- 公开：成员列表 ----------
@public_router.get("", response_model=list[MemberOut])
async def list_public_members(
    term: Literal["current", "former"] | None = Query(default=None),
    db: AsyncSession = Depends(get_db),
):
    """前端公开：按 display_order 升序，过滤归档项。"""
    stmt = select(Member).where(Member.archived.is_(False))
    if term:
        stmt = stmt.where(Member.term == MemberTerm(term))
    stmt = stmt.order_by(Member.display_order.asc(), Member.id.asc())
    result = await db.execute(stmt)
    return result.scalars().all()


# ---------- 管理：分页列表 ----------
@admin_router.get("", response_model=dict)
async def admin_list_members(
    term: Literal["current", "former"] | None = Query(default=None),
    archived: bool | None = Query(default=None),
    q: str | None = Query(default=None, max_length=64),
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=20, ge=1, le=100),
    user: dict = Depends(require_permission("members", "view")),
    db: AsyncSession = Depends(get_db),
):
    stmt = select(Member)
    if term:
        stmt = apply_eq(stmt, Member.term, MemberTerm(term))
    if archived is not None:
        stmt = apply_eq(stmt, Member.archived, archived)
    stmt = apply_search(stmt, [Member.name, Member.role_title, Member.bio], q)
    stmt = stmt.order_by(Member.display_order.asc(), Member.id.asc())
    return await paginate(db, stmt, page, page_size, out_model=MemberAdminOut)


# ---------- 管理：空白导入模板下载 ----------
# 注：必须定义在 /{mid} 之前，否则 "template" 会被当成成员 id
@admin_router.get("/template")
async def admin_download_member_template(
    user: dict = Depends(require_permission("members", "manage")),
):
    """下载空白 Excel 导入模板：表头复刻社团成员信息表，填完后走 /import 上传。"""
    headers = [
        ("职务", "role_title"),
        ("姓名", "name"),
        ("性别", "gender"),
        ("年级", "grade"),
        ("所在院系", "department"),
        ("专业及班级", "major_class"),
        ("联系电话", "phone"),
        ("微信号", "wechat"),
        ("政治面貌", "political_status"),
        ("是否为留学生", "is_intl_student"),
        ("届别", "term"),
    ]
    data = to_xlsx([], headers, text_columns={"phone"}, sheet_title="成员导入模板", add_index=False)
    return Response(
        content=data,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="members_template.xlsx"'},
    )


# ---------- 管理：Excel 批量导入 ----------
MEMBER_IMPORT_MAX_SIZE = 5 * 1024 * 1024  # 5MB
MEMBER_IMPORT_MAX_ROWS = 1000

# 表头别名 → 内部字段（兼容原信息表与模板两种表头写法）
_IMPORT_ALIASES = {
    "姓名": "name",
    "职务": "role_title",
    "性别": "gender",
    "年级": "grade",
    "年級": "grade",
    "所在院系": "department",
    "院系": "department",
    "学院": "department",
    "专业及班级": "major_class",
    "专业班级": "major_class",
    "专业": "major_class",
    "班级": "major_class",
    "联系电话": "phone",
    "电话": "phone",
    "手机号": "phone",
    "微信号": "wechat",
    "微信": "wechat",
    "政治面貌": "political_status",
    "是否为留学生": "is_intl_student",
    "留学生": "is_intl_student",
    "是否留学生": "is_intl_student",
    "届别": "term",
}

_FIELD_MAXLEN = {
    "name": 64, "role_title": 64, "gender": 8, "grade": 32,
    "department": 64, "major_class": 64, "phone": 32, "wechat": 64,
    "political_status": 32,
}


def _clean_cell(v) -> str | None:
    if v is None:
        return None
    s = str(v).strip().replace("\n", " ")
    return s or None


def _parse_term_cell(v) -> str | None:
    """届别单元格 → current/former；空为 None（新建默认 current，更新保持原值）；非法抛错。"""
    s = _clean_cell(v)
    if s is None:
        return None
    if s in ("现任", "current"):
        return "current"
    if s in ("往届", "former"):
        return "former"
    raise ValueError(f"届别只能填「现任」或「往届」，当前为「{s}」")


def _parse_intl_cell(v) -> bool | None:
    """留学生单元格 → True/False；空为 None（保持默认/原值）；非法抛错。"""
    s = _clean_cell(v)
    if s is None:
        return None
    if s in ("是", "yes", "YES", "true", "TRUE", "1"):
        return True
    if s in ("否", "no", "NO", "false", "FALSE", "0"):
        return False
    raise ValueError(f"是否为留学生只能填「是」或「否」，当前为「{s}」")


@admin_router.post("/import", response_model=MemberImportOut)
async def admin_import_members(
    request: Request,
    file: UploadFile = File(...),
    mode: Literal["skip", "update"] = Form(default="skip"),
    user: dict = Depends(require_permission("members", "manage")),
    db: AsyncSession = Depends(get_db),
):
    """Excel 批量导入成员：按姓名去重。

    - mode=skip：已存在的姓名整行跳过；mode=update：用非空单元格覆盖更新。
    - 表头行自动定位（含「姓名」的首行，最多向前找 15 行），兼容原信息表格式。
    - 新建成员默认届别现任、未归档；display_order 默认 0。
    """
    data = await file.read(MEMBER_IMPORT_MAX_SIZE + 1)
    if len(data) > MEMBER_IMPORT_MAX_SIZE:
        raise HTTPException(status_code=400, detail="文件大小不能超过 5MB")
    fname = (file.filename or "").lower()
    if not fname.endswith((".xlsx", ".xlsm")):
        raise HTTPException(status_code=400, detail="仅支持 .xlsx 格式（请先下载空白模板填写）")

    try:
        from openpyxl import load_workbook
    except ImportError:
        raise HTTPException(status_code=500, detail="服务端缺少 Excel 解析依赖")
    import io as _io

    try:
        wb = load_workbook(filename=_io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.active
    except Exception:
        raise HTTPException(status_code=400, detail="Excel 文件无法解析，请检查文件是否损坏")

    # 定位表头行
    header_row_idx: int | None = None
    col_map: dict[int, str] = {}
    for r in range(1, min(ws.max_row, 15) + 1):
        mapping: dict[int, str] = {}
        for c in range(1, ws.max_column + 1):
            label = _clean_cell(ws.cell(r, c).value)
            if label and label in _IMPORT_ALIASES:
                mapping[c] = _IMPORT_ALIASES[label]
        if "name" in mapping.values():
            header_row_idx = r
            col_map = mapping
            break
    if header_row_idx is None:
        raise HTTPException(status_code=400, detail="未找到含「姓名」的表头行，请使用下载的空白模板填写")

    existing = (await db.execute(select(Member))).scalars().all()
    by_name: dict[str, Member] = {m.name: m for m in existing}

    created = updated = skipped = 0
    errors: list[MemberImportError] = []
    seen_in_file: set[str] = set()
    n_data = 0

    for r in range(header_row_idx + 1, ws.max_row + 1):
        if n_data >= MEMBER_IMPORT_MAX_ROWS:
            errors.append(MemberImportError(row=r, name="", reason=f"超过单次最多 {MEMBER_IMPORT_MAX_ROWS} 行，剩余行未处理"))
            break
        cells = {field: _clean_cell(ws.cell(r, c).value) for c, field in col_map.items()}
        name = cells.get("name")
        if name is None:
            continue  # 纯空行跳过，不计数
        n_data += 1
        if len(name) > 64:
            errors.append(MemberImportError(row=r, name=name[:20], reason="姓名超过 64 字"))
            continue
        if name in seen_in_file:
            skipped += 1
            continue
        seen_in_file.add(name)

        try:
            term = _parse_term_cell(cells.get("term"))
            is_intl = _parse_intl_cell(cells.get("is_intl_student"))
        except ValueError as e:
            errors.append(MemberImportError(row=r, name=name, reason=str(e)))
            continue
        overlong = [k for k, v in cells.items() if k in _FIELD_MAXLEN and v and len(v) > _FIELD_MAXLEN[k]]
        if overlong:
            errors.append(MemberImportError(row=r, name=name, reason=f"字段过长：{','.join(overlong)}"))
            continue

        old = by_name.get(name)
        if old is not None and mode == "skip":
            skipped += 1
            continue
        payload = {
            "role_title": cells.get("role_title"),
            "gender": cells.get("gender"),
            "grade": cells.get("grade"),
            "department": cells.get("department"),
            "major_class": cells.get("major_class"),
            "phone": cells.get("phone"),
            "wechat": cells.get("wechat"),
            "political_status": cells.get("political_status"),
        }
        if old is not None:  # mode == update：仅非空覆盖；届别/留学生仅显式填写时覆盖
            for k, v in payload.items():
                if v is not None:
                    setattr(old, k, v)
            if term is not None:
                old.term = MemberTerm(term)
            if is_intl is not None:
                old.is_intl_student = is_intl
            updated += 1
        else:
            m = Member(
                name=name,
                term=MemberTerm(term) if term else MemberTerm.CURRENT,
                is_intl_student=is_intl if is_intl is not None else False,
                **payload,
            )
            db.add(m)
            by_name[name] = m
            created += 1

    _log(db, int(user["user_id"]), f"member.import_{mode}",
         f"created={created} updated={updated} skipped={skipped} errors={len(errors)}",
         get_client_ip(request))
    await db.commit()
    return MemberImportOut(total=n_data, created=created, updated=updated, skipped=skipped, errors=errors)


# ---------- 管理：批量修改届别 / 职务 ----------
@admin_router.post("/batch")
async def admin_batch_update_members(
    body: MemberBatchUpdate,
    request: Request,
    user: dict = Depends(require_permission("members", "manage")),
    db: AsyncSession = Depends(get_db),
):
    """批量勾选修改届别与职务：两项至少指定一项，为空的那项不改动。"""
    if body.term is None and body.role_title is None:
        raise HTTPException(status_code=400, detail="请至少指定届别或职务其中一项")
    result = await db.execute(select(Member).where(Member.id.in_(body.ids)))
    members = result.scalars().all()
    found = {m.id for m in members}
    updated = 0
    for m in members:
        if body.term is not None:
            m.term = MemberTerm(body.term)
        if body.role_title is not None:
            m.role_title = body.role_title
        updated += 1
    _log(db, int(user["user_id"]), "member.batch_update", ",".join(str(i) for i in body.ids), get_client_ip(request))
    await db.commit()
    return {"updated": updated, "not_found": sorted(set(body.ids) - found)}


# ---------- 管理：单条详情 ----------
@admin_router.get("/{mid}", response_model=MemberAdminOut)
async def admin_get_member(
    mid: int,
    user: dict = Depends(require_permission("members", "view")),
    db: AsyncSession = Depends(get_db),
):
    m = await db.get(Member, mid)
    if not m:
        raise HTTPException(status_code=404, detail="成员不存在")
    return m


# ---------- 管理：创建 ----------
@admin_router.post("", response_model=MemberAdminOut, status_code=status.HTTP_201_CREATED)
async def admin_create_member(
    req: MemberCreate,
    request: Request,
    user: dict = Depends(require_permission("members", "manage")),
    db: AsyncSession = Depends(get_db),
):
    m = Member(
        name=req.name,
        role_title=req.role_title,
        term=MemberTerm(req.term),
        bio=req.bio,
        avatar_url=req.avatar_url,
        display_order=req.display_order,
        archived=req.archived,
        gender=req.gender,
        grade=req.grade,
        department=req.department,
        major_class=req.major_class,
        phone=req.phone,
        wechat=req.wechat,
        political_status=req.political_status,
        is_intl_student=req.is_intl_student,
    )
    db.add(m)
    _log(db, int(user["user_id"]), "member.create", None, get_client_ip(request))
    await db.commit()
    await db.refresh(m)
    return m


# ---------- 管理：更新 ----------
@admin_router.put("/{mid}", response_model=MemberAdminOut)
async def admin_update_member(
    mid: int,
    req: MemberUpdate,
    request: Request,
    user: dict = Depends(require_permission("members", "manage")),
    db: AsyncSession = Depends(get_db),
):
    m = await db.get(Member, mid)
    if not m:
        raise HTTPException(status_code=404, detail="成员不存在")

    data = req.model_dump(exclude_unset=True)
    if "term" in data and data["term"] is not None:
        m.term = MemberTerm(data["term"])
    for k in ("name", "role_title", "bio", "avatar_url", "display_order", "archived",
                "gender", "grade", "department", "major_class", "phone", "wechat",
                "political_status", "is_intl_student"):
        if k in data:
            setattr(m, k, data[k])
    _log(db, int(user["user_id"]), "member.update", f"member:{mid}", get_client_ip(request))
    await db.commit()
    await db.refresh(m)
    return m


# ---------- 管理：删除 ----------
@admin_router.delete("/{mid}", status_code=status.HTTP_204_NO_CONTENT)
async def admin_delete_member(
    mid: int,
    request: Request,
    user: dict = Depends(require_permission("members", "manage")),
    db: AsyncSession = Depends(get_db),
):
    m = await db.get(Member, mid)
    if not m:
        raise HTTPException(status_code=404, detail="成员不存在")
    # 删除成员头像文件
    if m.avatar_url:
        try:
            old_path = (MEMBER_AVATAR_DIR / Path(m.avatar_url).name).resolve()
            if old_path.parent == MEMBER_AVATAR_DIR.resolve() and old_path.exists():
                old_path.unlink(missing_ok=True)
        except Exception:
            pass
    await db.delete(m)
    _log(db, int(user["user_id"]), "member.delete", f"member:{mid}", get_client_ip(request))
    await db.commit()


# ---------- 管理：头像上传 ----------
@admin_router.post("/avatar", response_model=AvatarUploadOut)
async def admin_upload_member_avatar(
    request: Request,
    file: UploadFile = File(...),
    user: dict = Depends(require_permission("members", "manage")),
    db: AsyncSession = Depends(get_db),
):
    """成员头像上传：jpg/png/webp，最大 5MB，自动压缩到最长边 512px。

    返回压缩后的图片 URL，前端直接写入 member.avatar_url 字段。
    """
    data = await file.read(MEMBER_AVATAR_MAX_SIZE + 1)
    if len(data) > MEMBER_AVATAR_MAX_SIZE:
        raise HTTPException(status_code=400, detail="文件大小不能超过 5MB")
    fmt = _sniff_image(data)
    if not fmt:
        raise HTTPException(status_code=400, detail="仅支持 JPG/PNG/WEBP 格式")

    # 压缩
    compressed, ext, mime = _compress_avatar(data, fmt)

    filename = f"member_{uuid.uuid4().hex}{ext}"
    MEMBER_AVATAR_DIR.mkdir(parents=True, exist_ok=True)
    (MEMBER_AVATAR_DIR / filename).write_bytes(compressed)
    avatar_url = f"/uploads/members/{filename}"

    # 记录到 media_files
    media = MediaFile(
        filename=filename,
        original_name=file.filename,
        mime=mime,
        size=len(compressed),
        storage_path=f"uploads/members/{filename}",
        uploader_id=int(user["user_id"]),
        category=MediaCategory.MEMBER,
    )
    db.add(media)
    _log(db, int(user["user_id"]), "member.avatar.upload", None, get_client_ip(request))
    await db.commit()
    return AvatarUploadOut(avatar_url=avatar_url)
