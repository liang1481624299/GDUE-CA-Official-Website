"""报名数据导出工具：CSV + Excel(.xlsx)。学号/手机号在 Excel 中强制为文本格式。"""
from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.cell.cell import WriteOnlyCell
from openpyxl.styles import Font

from app.db.models import Registration

HEADERS_CN = [
    ("序号", "index"),
    ("姓名", "name"),
    ("学号", "student_id"),
    ("学院", "college"),
    ("专业", "major"),
    ("区号", "phone_cc"),
    ("手机号", "phone_number"),
    ("邮箱", "email"),
    ("状态", "status"),
    ("备注", "remark"),
    ("提交时间", "submitted_at"),
    ("提交IP", "submit_ip"),
]


def _fmt(v) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d %H:%M:%S")
    if hasattr(v, "value"):  # Enum
        return v.value
    return str(v)


def registrations_to_csv(rows: list[Registration]) -> bytes:
    """导出 CSV（UTF-8 BOM + CRLF，Excel 直接打开不乱码）。"""
    # 简单 CSV 实现：字段内逗号/引号/换行转义
    def esc(s: str) -> str:
        if any(c in s for c in [",", '"', "\n", "\r"]):
            return '"' + s.replace('"', '""') + '"'
        return s

    lines = [",".join(h for h, _ in HEADERS_CN)]
    for i, r in enumerate(rows, 1):
        row = [str(i)]
        for _, key in HEADERS_CN[1:]:
            row.append(esc(_fmt(getattr(r, key))))
        lines.append(",".join(row))
    text = "\r\n".join(lines) + "\r\n"
    return ("\ufeff" + text).encode("utf-8")  # BOM 让 Excel 正确识别 UTF-8


def registrations_to_xlsx(rows: list[Registration]) -> bytes:
    """
    导出 Excel(.xlsx)。学号 / 手机号字段强制设置 number_format="@"（文本格式），
    避免 Excel 自动转为科学计数法或丢失前导零。
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "报名记录"

    # 表头
    header_font = Font(bold=True)
    for col, (label, _) in enumerate(HEADERS_CN, start=1):
        cell = ws.cell(row=1, column=col, value=label)
        cell.font = header_font

    # 数据行
    TEXT_COLUMNS = {"student_id", "phone_number"}  # 这两列强制文本
    for i, r in enumerate(rows, 1):
        for col, (_, key) in enumerate(HEADERS_CN, start=1):
            value = str(i) if key == "index" else _fmt(getattr(r, key))
            cell = ws.cell(row=i + 1, column=col, value=value)
            if key in TEXT_COLUMNS:
                cell.number_format = "@"  # 强制文本，不做任何数值解析

    # 列宽（粗略）
    widths = [6, 12, 14, 20, 18, 8, 16, 20, 12, 20, 20]
    for col, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = w

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
