"""安全合规回归测试：口令策略、登录锁定、CSRF、角色 / 内网访问控制、输入过滤、信息泄露。"""
from fastapi.testclient import TestClient

from tests.conftest import ADMIN_EMAIL, STRONG_ADMIN_PASSWORD, csrf, make_client

ADMIN_USERNAME = ADMIN_EMAIL.split("@")[0]


def _login(client: TestClient, username: str, password: str, **extra):
    return client.post(
        "/api/auth/login",
        json={"username": username, "password": password, **extra},
        headers=csrf(client),
    )


def _admin_client(app) -> tuple[TestClient, dict]:
    """登录超管并完成首次改密，返回 (client, csrf 头)。"""
    client = make_client(app)
    with client:
        pass
    res = _login(client, ADMIN_USERNAME, STRONG_ADMIN_PASSWORD)
    if res.status_code == 401:  # 其他用例已改过密码
        res = _login(client, ADMIN_USERNAME, STRONG_ADMIN_PASSWORD + "2")
    assert res.status_code == 200, res.text
    headers = {"X-CSRF-Token": res.json()["csrf_token"]}
    if res.json()["must_change_password"]:
        r = client.post("/api/auth/change-password", headers=headers, json={
            "old_password": STRONG_ADMIN_PASSWORD, "new_password": STRONG_ADMIN_PASSWORD + "2",
        })
        assert r.status_code == 200, r.text
    return client, headers


# ---------- 密码存储与强口令 ----------
def test_password_hashed_with_bcrypt(app):
    import asyncio
    from sqlalchemy import select
    from app.db.models import User
    from app.db.session import async_session

    with make_client(app):
        pass

    async def fetch():
        async with async_session() as s:
            return (await s.execute(select(User).where(User.email == ADMIN_EMAIL))).scalar_one()
    user = asyncio.run(fetch())
    assert user.password_hash.startswith("$2b$12$")
    assert STRONG_ADMIN_PASSWORD not in user.password_hash
    assert user.must_change_password is True


def test_weak_passwords_rejected():
    from app.core.password_policy import check_password_strength
    for weak in ["admin", "12345678", "Password1", "aaaaBBBB1111", "Abcd1234!xyz",
                 "Qwer!9876zz", "admin@123", "short1A!"]:
        assert check_password_strength(weak) is not None, weak
    assert check_password_strength("zhangsan#2025Ok", username="zhangsan") is not None
    assert check_password_strength("Gd#Ca-Str0ng!Pwd") is None


def test_change_password_enforces_policy(app):
    client, headers = _admin_client(app)
    r = client.post("/api/auth/change-password", headers=headers, json={
        "old_password": STRONG_ADMIN_PASSWORD + "2", "new_password": "admin123",
    })
    assert r.status_code == 400


# ---------- 会话与 CSRF ----------
def test_login_sets_httponly_cookie_and_no_token_in_body(app):
    client = make_client(app)
    res = _login(client, ADMIN_USERNAME, STRONG_ADMIN_PASSWORD + "2")
    if res.status_code == 401:
        res = _login(client, ADMIN_USERNAME, STRONG_ADMIN_PASSWORD)
    assert res.status_code == 200
    assert "access_token" not in res.json()
    set_cookie = ",".join(res.headers.get_list("set-cookie"))
    assert "gdueca_session=" in set_cookie
    assert "HttpOnly" in set_cookie and "SameSite=strict" in set_cookie


def test_write_without_csrf_token_rejected(app):
    client, headers = _admin_client(app)
    r = client.put("/api/auth/profile", json={"real_name": "测试"})
    assert r.status_code == 403 and r.json()["code"] == "csrf_failed"
    r = client.put("/api/auth/profile", json={"real_name": "测试"},
                   headers={"X-CSRF-Token": "forged-token-value-forged-token-value"})
    assert r.status_code == 403
    r = client.put("/api/auth/profile", json={"real_name": "测试"},
                   headers={**headers, "Origin": "https://evil.example.com"})
    assert r.status_code == 403
    r = client.put("/api/auth/profile", json={"real_name": "测试"}, headers=headers)
    assert r.status_code == 200, r.text


def test_public_form_requires_csrf(app):
    client = make_client(app)
    payload = {"contact_email": "a@b.com", "description": "页面报错"}
    assert client.post("/api/bugs", json=payload).status_code == 403
    assert client.post("/api/bugs", json=payload, headers=csrf(client)).status_code == 201


def test_logout_revokes_session(app):
    client, headers = _admin_client(app)
    assert client.get("/api/auth/me").status_code == 200
    cookie = client.cookies.get("gdueca_session")
    assert client.post("/api/auth/logout", headers=headers).status_code == 200
    # 即使攻击者保留了旧 Cookie，会话也已在服务端撤销
    other = make_client(app)
    other.cookies.set("gdueca_session", cookie)
    assert other.get("/api/auth/me").status_code == 401


# ---------- 身份鉴别：登录失败锁定 ----------
def test_account_locks_after_repeated_failures(app):
    client, headers = _admin_client(app)
    r = client.post("/api/auth/users", headers=headers, json={
        "username": "locktest", "email": "lock@gdue-ca.cn", "password": "Zq#Vault-2025x",
        "role": "member", "student_id": "20250001", "real_name": "锁定测试", "phone": "+8613800000001",
    })
    assert r.status_code == 201, r.text
    anon = make_client(app)
    for _ in range(3):
        assert _login(anon, "locktest", "wrong-password").status_code == 401
    # 已锁定：即使密码正确也拒绝
    assert _login(anon, "locktest", "Zq#Vault-2025x").status_code == 423


# ---------- 访问控制：禁止匿名、角色区分、内网限制 ----------
def test_admin_endpoints_reject_anonymous(app):
    client = make_client(app)
    for path in ["/api/auth/users", "/api/registrations", "/api/bugs",
                 "/api/system/settings/admin", "/api/admin-stats/access", "/api/auth/me"]:
        assert client.get(path).status_code == 401, path


def test_member_cannot_access_admin(app):
    client, headers = _admin_client(app)
    r = client.post("/api/auth/users", headers=headers, json={
        "username": "member1", "email": "member1@gdue-ca.cn", "password": "Memb#er-2025xy",
        "role": "member", "student_id": "20250002", "real_name": "普通成员", "phone": "+8613800000002",
    })
    assert r.status_code == 201, r.text
    m = make_client(app, ip="8.8.8.8")  # 普通成员允许外网登录
    res = _login(m, "member1", "Memb#er-2025xy")
    assert res.status_code == 200, res.text
    h = {"X-CSRF-Token": res.json()["csrf_token"]}
    m.post("/api/auth/change-password", headers=h, json={
        "old_password": "Memb#er-2025xy", "new_password": "Memb#er-2025xyZ"})
    assert m.get("/api/auth/profile").status_code == 200
    for path in ["/api/registrations", "/api/bugs", "/api/auth/users", "/api/system/settings/admin"]:
        assert m.get(path).status_code == 403, path
    # 不能给自己提权
    me = m.get("/api/auth/me").json()
    assert m.put(f"/api/auth/users/{me['id']}", headers=h, json={"role": "super_admin"}).status_code == 403


def test_admin_login_denied_from_public_network(app):
    client = make_client(app, ip="8.8.8.8")
    res = _login(client, ADMIN_USERNAME, STRONG_ADMIN_PASSWORD + "2")
    assert res.status_code == 403


def test_admin_session_unusable_from_public_network(app):
    client, _ = _admin_client(app)
    cookie = client.cookies.get("gdueca_session")
    outside = make_client(app, ip="203.0.113.9")
    outside.cookies.set("gdueca_session", cookie)
    assert outside.get("/api/auth/users").status_code == 403


def test_spoofed_xff_from_untrusted_peer_ignored(app):
    """非可信代理直连时伪造 X-Forwarded-For: 127.0.0.1 不能冒充内网。"""
    from fastapi.testclient import TestClient
    client = TestClient(app, client=("8.8.4.4", 40000))
    client.headers["X-Forwarded-For"] = "127.0.0.1"
    res = _login(client, ADMIN_USERNAME, STRONG_ADMIN_PASSWORD + "2")
    assert res.status_code == 403


def test_recover_is_intranet_only(app):
    outside = make_client(app, ip="8.8.8.8")
    assert outside.get("/api/auth/security-question").status_code == 403


# ---------- XSS：输入过滤 ----------
def test_xss_payloads_rejected_in_plain_fields(app):
    client = make_client(app)
    base = {
        "name": "<script>alert(1)</script>", "student_id": "20250003", "college": "计算机学院",
        "major": "软件工程", "phone_cc": "+86", "phone_number": "13800000003",
        "position": "技术部", "introduction": "你好",
    }
    r = client.post("/api/registrations/club", json=base, headers=csrf(client))
    assert r.status_code == 422
    r = client.post("/api/registrations/club", headers=csrf(client),
                    json={**base, "name": "张三", "college": '<img src=x onerror=alert(1)>'})
    assert r.status_code == 422
    r = client.post("/api/registrations/club", headers=csrf(client),
                    json={**base, "name": "张三", "student_id": "2025\"><svg>"})
    assert r.status_code == 422


def test_activity_cover_url_rejects_javascript_scheme(app):
    client, headers = _admin_client(app)
    r = client.post("/api/activities", headers=headers, json={
        "title": "测试活动", "content": "内容", "cover_url": "javascript:alert(1)",
    })
    assert r.status_code == 422


def test_api_security_headers(app):
    client = make_client(app)
    r = client.get("/api/system/settings")
    assert r.headers["x-content-type-options"] == "nosniff"
    assert r.headers["x-frame-options"] == "DENY"
    assert "default-src 'none'" in r.headers["content-security-policy"]
    assert r.headers["cache-control"] == "no-store"


# ---------- 信息泄露 ----------
def test_public_settings_hide_security_config(app):
    data = make_client(app).get("/api/system/settings").json()
    for key in ("ip_blacklist", "allowed_hosts", "cors_origins", "network_listen_ip"):
        assert key not in data


def test_draft_activity_hidden_from_public(app):
    client, headers = _admin_client(app)
    r = client.post("/api/activities", headers=headers,
                    json={"title": "草稿活动", "content": "x", "category": "讲座", "status": "draft"})
    assert r.status_code == 201, r.text
    aid = r.json()["id"]
    anon = make_client(app)
    assert anon.get(f"/api/activities/{aid}").status_code == 404
    assert all(a["id"] != aid for a in anon.get("/api/activities?status_filter=all").json())
    assert any(a["id"] == aid for a in client.get("/api/activities?status_filter=all").json())


def test_export_neutralizes_formulas():
    from types import SimpleNamespace
    from app.utils.export import registrations_to_csv, registrations_to_xlsx
    row = SimpleNamespace(name="=HYPERLINK(\"http://evil\")", student_id="1", college="c", major="m",
                          phone_cc="+86", phone_number="138", email=None, status="pending",
                          remark="@SUM(1)", submitted_at=None, submit_ip=None)
    text = registrations_to_csv([row]).decode("utf-8-sig")
    assert "'=HYPERLINK" in text and "'@SUM" in text and ",+86," in text
    from io import BytesIO
    from openpyxl import load_workbook
    ws = load_workbook(BytesIO(registrations_to_xlsx([row]))).active
    assert ws.cell(row=2, column=2).data_type == "s"


# ---------- 个人信息最小化：公开提交接口只返回回执 ----------
def test_public_submit_returns_receipt_only(app):
    client = make_client(app)
    r = client.post("/api/registrations/club", headers=csrf(client), json={
        "name": "张三", "student_id": "20250009", "college": "计算机学院", "major": "软件工程",
        "phone_cc": "+86", "phone_number": "13800000009", "email": "z@qq.com",
        "position": "技术部", "introduction": "你好",
    })
    assert r.status_code == 201, r.text
    assert set(r.json()) == {"receipt_code", "submitted_at"}
    r = client.post("/api/bugs", headers=csrf(client),
                    json={"contact_email": "a@b.com", "description": "页面报错"})
    assert r.status_code == 201, r.text
    assert set(r.json()) == {"receipt_code", "submitted_at"}
